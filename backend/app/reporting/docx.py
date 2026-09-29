"""Editable Word export built from the same report model (python-docx)."""
from __future__ import annotations

import io
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _kv_table(doc: Document, items: list[list[str]]) -> None:
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in items:
        row = t.add_row().cells
        row[0].text = str(k)
        row[1].text = "—" if v in (None, "") else str(v)
        row[0].paragraphs[0].runs[0].bold = True if row[0].paragraphs[0].runs else True
    doc.add_paragraph()


def _grid_table(doc: Document, columns: list[str], rows: list[list]) -> None:
    t = doc.add_table(rows=1, cols=len(columns))
    t.style = "Table Grid"
    for i, c in enumerate(columns):
        t.rows[0].cells[i].text = str(c)
        for p in t.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row[: len(columns)]):
            cells[i].text = "—" if c in (None, "") else str(c)
    doc.add_paragraph()


def build_docx(model: dict) -> bytes:
    doc = Document()

    h = doc.add_heading(model["laboratory"]["name"], level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("TYPE EVALUATION REPORT — OIML R 76")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    no = doc.add_paragraph(model["report_no"])
    no.alignment = WD_ALIGN_PARAGRAPH.CENTER
    no.runs[0].bold = True

    doc.add_heading("1. General information concerning the type", level=1)
    inst = model["instrument"]
    _kv_table(doc, [
        ["Application no.", inst["application_no"]],
        ["Type designation", inst["type_designation"]],
        ["Manufacturer", inst["manufacturer"]],
        ["Category", inst["category"]],
        ["Accuracy class", inst["accuracy_class"]],
        ["Indication type", inst["indication_type"]],
        ["Display", inst["display_kind"]],
        ["Range kind", inst["range_kind"]],
        ["Min", inst["min_capacity"]],
        ["Lim (Max safe load)", inst["max_safe_load"]],
        ["T = + / T = −", f"{inst['tare_plus']} / {inst['tare_minus']}"],
        ["Power supply", inst["power_supply"]],
        ["Temperature limits", inst["temp_limits"]],
        ["Printer", inst["printer"]],
        ["Software version", inst["software_version"]],
        ["Submitted serial no.", inst["submitted_serial"]],
    ])

    doc.add_heading("2. Weighing ranges", level=1)
    _grid_table(doc, ["Range", "e", "d", "Max", "n = Max/e"], model["ranges"])

    doc.add_heading("3. Test equipment (traceability)", level=1)
    if model["equipment"]:
        _grid_table(doc, ["Kind", "Name", "Serial", "Class/U", "Certificate"],
                    model["equipment"])

    doc.add_heading("4. Test conditions & rule set", level=1)
    c = model["conditions"]
    _kv_table(doc, [
        ["Purpose", c["purpose"]], ["MPE context", c["mpe_context"]],
        ["Auto zero device", c["auto_zero_status"]],
        ["Resolution during test", c["resolution_during_test"]],
        ["Evaluation period", c["period"]],
        ["Rule set", f"{c['ruleset_id']} (sha256 {c['ruleset_sha256'][:16]}…)"],
    ])

    for page in model["pages"]:
        doc.add_page_break()
        title = f"Form {page['form_no']}"
        doc.add_heading(f"{title} — verdict: {page['verdict']}", level=1)
        for b in page["blocks"]:
            if b["type"] == "kv":
                doc.add_heading(b["title"], level=2)
                _kv_table(doc, b["items"])
            elif b["type"] == "table":
                doc.add_heading(b["title"], level=2)
                _grid_table(doc, b["columns"], b["rows"])
            elif b["type"] == "note":
                doc.add_paragraph(b["text"])
            elif b["type"] == "json":
                doc.add_paragraph(json.dumps(b["data"], indent=2, default=str))
        if page["warnings"]:
            doc.add_heading("Warnings", level=2)
            for w in page["warnings"]:
                doc.add_paragraph(f"⚠ {w}", style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Conformity summary", level=1)
    _grid_table(doc, ["Test", "Verdict"], model["conformity"])
    p = doc.add_paragraph(f"OVERALL OUTCOME: {model['outcome']}")
    p.runs[0].bold = True

    doc.add_heading("Approval", level=1)
    _kv_table(doc, [
        ["Observer", model["observer"]],
        ["Approver", model["approver"] or "Pending review"],
        ["Approved at", model["approved_at"] or "—"],
        ["Approval hash (SHA-256)", model["approval_hash"] or "—"],
    ])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()